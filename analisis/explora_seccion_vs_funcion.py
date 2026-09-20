# -*- coding: utf-8 -*-
"""Que 'trabajo que realiza' hay detras de cada seccion de la gestoria.
Cruza (en memoria, por nombre) los bloques por seccion del RESUMEN de agosto
2026 con el 'Listado de personal por funciones'. Solo se imprimen recuentos."""
import collections
import re
import unicodedata

import openpyxl

LISTADO = r"Z:\LABORAL\2026-01-27 Listado de personal por funciones.xlsx"
RESUMENES = {
    "Razo": r"Z:\LABORAL\_NOMINAS\NOMINAS\2026\08 Agosto 2026\TR - RESUMEN NOMINA.xlsx",
    "Agetrans": r"Z:\LABORAL\_NOMINAS\NOMINAS\2026\08 Agosto 2026\AGETRANS - RESUMEN NOMINA.xlsx",
}
RUIDO = {"DE", "DEL", "LA", "LAS", "LOS", "Y"}


def toks(s):
    s = unicodedata.normalize("NFKD", str(s or "")).encode("ascii", "ignore").decode().upper()
    return {t for t in re.sub(r"[^A-Z ]", " ", s).split() if t not in RUIDO and len(t) > 1}


def jacc(a, b):
    return len(a & b) / len(a | b) if a and b else 0.0


# ---- listado: nombre -> (empresa, trabajo)
wb = openpyxl.load_workbook(LISTADO, read_only=True, data_only=True)
lista = []
empresa = None
for f in wb["Hoja1"].iter_rows(min_row=2, values_only=True):
    a, d = f[0], (f[3] if len(f) > 3 else None)
    if a and not any(f[1:4]):
        empresa = "Agetrans" if "AGETRANS" in str(a).upper() else "Razo"
        continue
    if a and d:
        lista.append((empresa, toks(a), str(d).strip()))
print("listado: %d personas con funcion; por empresa: %s" % (len(lista), collections.Counter(e for e, _, _ in lista)))

for emp, ruta in RESUMENES.items():
    wb = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
    filas = [list(f) for f in wb["Detalle"].iter_rows(values_only=True)]
    tot = [list(f) for f in wb["Totales"].iter_rows(values_only=True)]
    # totales por seccion (codigo -> TOTAL)
    ini = next(i for i, f in enumerate(tot) if f and isinstance(f[0], str) and f[0].upper().startswith("TOTALES SECCIONES"))
    cab = next(i for i in range(ini, ini + 6) if any(isinstance(v, (int, float)) for v in tot[i][2:]))
    fila_total = next(f for f in tot[cab + 1:] if f and isinstance(f[0], str) and f[0].strip() == "TOTAL")
    cod_por_col = {j: v for j, v in enumerate(tot[cab]) if isinstance(v, (int, float)) and j >= 3}
    total_por_codigo = {int(cod_por_col[j]): fila_total[j] for j in cod_por_col}
    # bloques del detalle
    inicios = [i for i, f in enumerate(filas) if len(f) > 1 and isinstance(f[1], str) and f[1].strip().startswith("Cód. sección")]
    print("\n##### %s: %d bloques de seccion en Detalle; TOTAL por codigo (Totales): %s" % (
        emp, len(inicios), {k: round(v, 2) for k, v in total_por_codigo.items()}))
    for n, i0 in enumerate(inicios):
        fin = inicios[n + 1] if n + 1 < len(inicios) else len(filas)
        bloque = filas[i0:fin]
        # total del bloque = columna B en la fila de etiqueta 'TOTAL'
        fila_t = next((f for f in bloque if isinstance(f[0], str) and f[0].strip() == "TOTAL"), None)
        total_bloque = fila_t[1] if fila_t else None
        codigo = next((k for k, v in total_por_codigo.items() if total_bloque is not None and abs(v - total_bloque) < 0.02), None)
        # empleados: columnas desde la 2 con codigo numerico en la fila i0+2
        cods = bloque[2] if len(bloque) > 2 else []
        emps = []
        for j in range(2, len(cods)):
            if cods[j] in (None, ""):
                continue
            partes = [bloque[k][j] for k in (3, 4, 5) if len(bloque) > k and len(bloque[k]) > j and bloque[k][j]]
            emps.append(toks(" ".join(str(p) for p in partes)))
        cuenta = collections.Counter()
        sin_match = 0
        for nombre in emps:
            mejor, sc = None, 0.0
            for e, t, trabajo in lista:
                if e != emp:
                    continue
                s = jacc(nombre, t)
                if s > sc:
                    mejor, sc = trabajo, s
            if mejor and sc >= 0.5:
                cuenta[mejor] += 1
            else:
                sin_match += 1
        print("  bloque %d -> codigo seccion %s | total=%s | empleados=%d | sin casar con listado=%d | funciones=%s" % (
            n + 1, codigo, None if total_bloque is None else round(total_bloque, 2), len(emps), sin_match, dict(cuenta)))
