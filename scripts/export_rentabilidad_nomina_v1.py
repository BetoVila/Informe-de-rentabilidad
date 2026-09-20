# -*- coding: utf-8 -*-
"""Lectura de los RESUMEN DE NOMINA mensuales de la gestoria (Z:\\LABORAL\\_NOMINAS).

Un fichero por empresa y mes, en .xls o .xlsx, con nombres muy variables.
Hoja 'Totales' -> 'TOTALES SECCIONES': una columna por seccion de la gestoria
(=el trabajo que realiza cada persona: conductor hormigonera, nacional,
banera, administracion, taller...). Hoja 'Detalle': una columna por persona.

COSTE DE EMPRESA EXACTO: fila TOTAL = TOTAL DEVENGOS + TOTAL COSTE S.S. (se
comprueba fichero a fichero; nada se recalcula ni se estima).

DOS SALIDAS, porque los sueldos con nombre son datos personales:
  --output          agregados por (empresa, mes, seccion), SIN nombres. Los
                    grupos de menos de 3 personas se pliegan en el mayor de su
                    empresa y mes, para que un total no sea el sueldo de UNA
                    persona. Va al informe que abre cualquiera con acceso a la
                    carpeta.
  --output-detail   coste por persona y mes, CON nombre. Solo para la capa
                    protegida con clave; vive unicamente en la carpeta privada
                    de trabajo.

DUPLICADOS: si una misma (empresa, mes) sale en varios ficheros -carpeta
equivocada, 'Copia de...', nomina reprocesada- manda el de fecha mas reciente
y las demas versiones quedan anotadas, con sus totales, para revisarlas.

SOLO LECTURA. Los ficheros de la gestoria no se tocan.
"""
import argparse
import collections
import datetime as dt
import glob
import json
import os
import re
import sys
import unicodedata

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(os.path.dirname(HERE), "vendor"))   # openpyxl, et_xmlfile, xlrd (Python puro)
import openpyxl  # noqa: E402
import xlrd  # noqa: E402

MESES = {"enero": 1, "febrero": 2, "marzo": 3, "abril": 4, "mayo": 5, "junio": 6, "julio": 7, "agosto": 8,
         "septiembre": 9, "setiembre": 9, "octubre": 10, "noviembre": 11, "diciembre": 12}
MIN_PERSONAS = 3          # por debajo, el grupo se pliega (privacidad)
CONCEPTOS = {"TOTAL": "cost", "TOTAL DEVENGOS": "devengos", "TOTAL COSTE S.S.": "ss",
             "040 DIETAS": "dietas", "035 H.EXTRAS": "extras"}


def norm(s):
    s = unicodedata.normalize("NFKD", str(s or "")).encode("ascii", "ignore").decode().upper()
    return re.sub(r"\s+", " ", s).strip()


def num(v):
    return float(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else None


def leer_hojas(ruta):
    """{hoja: filas}; celda vacia = None (xlrd devuelve '' y openpyxl None)."""
    hojas = {}
    if ruta.lower().endswith(".xlsx"):
        wb = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
        try:
            for n in wb.sheetnames:
                hojas[n] = [list(f) for f in wb[n].iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        wb = xlrd.open_workbook(ruta, on_demand=True)
        try:
            for n in wb.sheet_names():
                h = wb.sheet_by_name(n)
                hojas[n] = [[(None if v == "" else v) for v in h.row_values(i)] for i in range(h.nrows)]
                wb.unload_sheet(n)
        finally:
            wb.release_resources()
    return hojas


def empresa_de(texto):
    t = norm(texto)
    if "AGETRANS" in t or "B15938442" in t or re.match(r"^540\b", t):
        return "Agetrans"
    if "RAZO" in t or "B15226095" in t or re.match(r"^27180\b", t):
        return "Razo"
    return None


def cabecera(filas):
    emp = periodo = None
    for f in filas[:14]:
        a = norm(f[0]) if f and f[0] is not None else ""
        b = str(f[1]).strip() if len(f) > 1 and f[1] is not None else ""
        if a == "EMPRESA":
            emp = empresa_de(b)
        elif a == "PROCESO":
            m = re.search(r"([A-Za-z]+)\s+de[l]?\s+(\d{4})", unicodedata.normalize("NFKD", b).encode("ascii", "ignore").decode())
            if m and m.group(1).lower() in MESES:
                periodo = "%s-%02d" % (m.group(2), MESES[m.group(1).lower()])
    return emp, periodo


def parse_totales(filas):
    """{codigo_seccion: {cost,devengos,ss,dietas,extras}} y total empresa."""
    ini = next((i for i, f in enumerate(filas) if f and isinstance(f[0], str) and f[0].strip().upper().startswith("TOTALES SECCIONES")), None)
    if ini is None:
        raise ValueError("sin bloque TOTALES SECCIONES")
    cab = next((i for i in range(ini + 1, min(ini + 8, len(filas))) if len(filas[i]) > 1 and isinstance(filas[i][1], str) and filas[i][1].strip().lower() == "totales"), None)
    if cab is None:
        raise ValueError("sin fila de cabecera de secciones")
    cols, visto_num = {}, False
    for j in range(2, len(filas[cab])):
        v = filas[cab][j]
        if num(v) is not None:
            cols[j] = int(v)
            visto_num = True
        elif v is None and not visto_num:
            cols[j] = 0                               # columna sin cabecera = «sin seccion»
    por_seccion = {c: {} for c in cols.values()}
    total_empresa = {}
    for f in filas[cab + 1:]:
        if not f or not isinstance(f[0], str):
            continue
        clave = CONCEPTOS.get(f[0].strip())
        if not clave:
            continue
        total = num(f[1] if len(f) > 1 else None) or 0.0
        suma = 0.0
        for j, cod in cols.items():
            v = num(f[j] if len(f) > j else None) or 0.0
            por_seccion[cod][clave] = v
            suma += v
        total_empresa[clave] = total
        if abs(suma - total) > 0.05:
            raise ValueError("la suma de secciones no cuadra en %s: %.2f frente a %.2f" % (f[0].strip(), suma, total))
    for k in ("cost", "devengos", "ss"):
        if k not in total_empresa:
            raise ValueError("falta la fila " + k)
    if abs(total_empresa["devengos"] + total_empresa["ss"] - total_empresa["cost"]) > 0.05:
        raise ValueError("TOTAL != DEVENGOS + COSTE S.S.")
    return por_seccion, total_empresa


def parse_detalle(filas, por_seccion):
    """Personas por bloque de seccion. Devuelve lista de dicts (con nombre)."""
    inicios = [i for i, f in enumerate(filas) if len(f) > 1 and isinstance(f[1], str) and f[1].strip().startswith("Cód. sección")]
    if not inicios:
        raise ValueError("sin bloques de seccion en Detalle")
    personas = []
    codigos_usados = set()
    for n, i0 in enumerate(inicios):
        fin = inicios[n + 1] if n + 1 < len(inicios) else len(filas)
        bloque = filas[i0:fin]
        etiquetas = {}
        for k, f in enumerate(bloque):
            if f and isinstance(f[0], str):
                etiquetas.setdefault(f[0].strip(), k)
        if "TOTAL" not in etiquetas:
            raise ValueError("bloque %d sin fila TOTAL" % (n + 1))
        tot_bloque = num(bloque[etiquetas["TOTAL"]][1]) or 0.0
        # seccion del bloque = la cuyo TOTAL coincide (el numero entre parentesis es un ordinal, no el codigo)
        cand = [(abs(v.get("cost", 0.0) - tot_bloque), c) for c, v in por_seccion.items() if c not in codigos_usados]
        cand.sort()
        if not cand or cand[0][0] > 0.02:
            raise ValueError("el bloque %d no coincide con ninguna seccion de Totales" % (n + 1))
        codigo = cand[0][1]
        codigos_usados.add(codigo)
        fila_cod = bloque[2] if len(bloque) > 2 else []
        for j in range(2, len(fila_cod)):
            if fila_cod[j] in (None, ""):
                continue
            partes = [bloque[k][j] for k in (3, 4, 5) if len(bloque) > k and len(bloque[k]) > j and bloque[k][j] not in (None, "")]
            nombre_pila = str(partes[0]).strip() if partes else ""
            apellidos = " ".join(str(p).strip() for p in partes[1:])
            rec = {"seccion": codigo, "codigo": str(fila_cod[j]).strip(),
                   "nombre": ("%s, %s" % (apellidos, nombre_pila)) if apellidos else nombre_pila}
            for etq, clave in CONCEPTOS.items():
                if etq in etiquetas and len(bloque[etiquetas[etq]]) > j:
                    rec[clave] = num(bloque[etiquetas[etq]][j]) or 0.0
                else:
                    rec[clave] = 0.0
            personas.append(rec)
        suma = sum(p["cost"] for p in personas if p["seccion"] == codigo)
        if abs(suma - tot_bloque) > 0.05:
            raise ValueError("las personas del bloque %d suman %.2f y el bloque dice %.2f" % (n + 1, suma, tot_bloque))
    return personas


def plegar(filas):
    """Grupos de <MIN_PERSONAS -> se suman al mayor de su empresa y mes."""
    por_grupo = collections.defaultdict(list)
    for r in filas:
        por_grupo[(r["company"], r["period"])].append(r)
    salida = []
    for (_, _), rs in por_grupo.items():
        mayor = max(rs, key=lambda r: r["cost"])
        for r in rs:
            if r is mayor or (r["people"] is not None and r["people"] >= MIN_PERSONAS):
                continue
            for k in ("cost", "devengos", "ss", "dietas", "extras"):
                mayor[k] = round(mayor[k] + r[k], 2)
            mayor["people"] = (mayor["people"] or 0) + (r["people"] or 0)
            mayor["plegado"] += 1 + r["plegado"]
            r["_borrar"] = True
        salida.extend(r for r in rs if not r.get("_borrar"))
    return salida


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--root", required=True, help=r"Carpeta Z:\LABORAL\_NOMINAS")
    ap.add_argument("--output", required=True)
    ap.add_argument("--output-detail", default="")
    ap.add_argument("--from-date", default="2025-01-01")
    ap.add_argument("--to-date", default=(dt.date.today() - dt.timedelta(days=1)).isoformat())
    args = ap.parse_args()
    desde_mes, hasta_mes = args.from_date[:7], args.to_date[:7]

    meta = {"desde": args.from_date, "hasta": args.to_date, "read_at": dt.datetime.now().astimezone().isoformat(),
            "disponible": False, "minPersonasGrupo": MIN_PERSONAS}
    if not os.path.isdir(args.root):
        meta["aviso"] = "No se encuentra la carpeta de nominas: " + args.root
        json.dump({"metadata": meta, "rows": [], "files": [], "versions": [], "errors": []},
                  open(args.output, "w", encoding="utf-8"), ensure_ascii=False, separators=(",", ":"))
        print(json.dumps({"disponible": False, "aviso": meta["aviso"]}, ensure_ascii=False))
        return 0

    candidatos = []
    for patron in (os.path.join(args.root, "NOMINAS", "*", "*", "*"), os.path.join(args.root, "datos nominas", "*")):
        for ruta in glob.glob(patron):
            n = os.path.basename(ruta)
            if ruta.lower().endswith((".xls", ".xlsx")) and "resumen" in n.lower() and not n.startswith("~$"):
                candidatos.append(ruta)

    leidos, errores = [], []
    for ruta in sorted(candidatos):
        try:
            hojas = leer_hojas(ruta)
            h_tot = next((h for n, h in hojas.items() if "total" in n.lower()), None)
            h_det = next((h for n, h in hojas.items() if "detalle" in n.lower()), None)
            if h_tot is None:
                raise ValueError("sin hoja Totales")
            emp, periodo = cabecera(h_tot)
            if not emp or not periodo:
                raise ValueError("no se reconoce empresa/periodo (%s, %s)" % (emp, periodo))
            if not (desde_mes <= periodo <= hasta_mes):
                continue
            por_seccion, total_empresa = parse_totales(h_tot)
            personas = None
            aviso_detalle = None
            if h_det is not None:
                try:
                    personas = parse_detalle(h_det, por_seccion)
                except ValueError as e:
                    aviso_detalle = str(e)
            leidos.append({"ruta": ruta, "mtime": os.path.getmtime(ruta), "company": emp, "period": periodo,
                           "por_seccion": por_seccion, "total": total_empresa, "personas": personas, "aviso_detalle": aviso_detalle})
        except Exception as e:  # noqa: BLE001 - un fichero malo no tumba a los demas
            errores.append({"ruta": ruta, "error": "%s: %s" % (type(e).__name__, str(e)[:160])})

    grupos = collections.defaultdict(list)
    for r in leidos:
        grupos[(r["company"], r["period"])].append(r)
    elegidos, versiones = [], []
    for k, rs in sorted(grupos.items()):
        rs.sort(key=lambda r: r["mtime"])
        elegido = rs[-1]
        elegidos.append(elegido)
        if len(rs) > 1:
            versiones.append({"company": k[0], "period": k[1], "usada": os.path.relpath(elegido["ruta"], args.root),
                              "difieren": any(abs(r["total"]["cost"] - elegido["total"]["cost"]) > 0.005 for r in rs),
                              "versiones": [{"fichero": os.path.relpath(r["ruta"], args.root),
                                             "modificado": dt.datetime.fromtimestamp(r["mtime"]).isoformat(timespec="minutes"),
                                             "coste": round(r["total"]["cost"], 2)} for r in rs]})

    filas, detalle = [], []
    for r in elegidos:
        det = r["personas"]
        for cod, v in sorted(r["por_seccion"].items()):
            gente = None if det is None else sum(1 for p in det if p["seccion"] == cod)
            if gente == 0 and abs(v.get("cost", 0.0)) < 0.005:
                continue
            filas.append({"company": r["company"], "period": r["period"], "section": cod, "people": gente,
                          "cost": round(v.get("cost", 0.0), 2), "devengos": round(v.get("devengos", 0.0), 2),
                          "ss": round(v.get("ss", 0.0), 2), "dietas": round(v.get("dietas", 0.0), 2),
                          "extras": round(v.get("extras", 0.0), 2), "plegado": 0})
        for p in det or []:
            detalle.append({"company": r["company"], "period": r["period"], **{k: (round(p[k], 2) if isinstance(p[k], float) else p[k]) for k in p}})
    filas = plegar(filas)
    filas.sort(key=lambda r: (r["company"], r["period"], r["section"]))

    periodos = collections.defaultdict(list)
    for r in elegidos:
        periodos[r["company"]].append(r["period"])
    meta.update({"disponible": True, "ficheros": len(leidos), "elegidos": len(elegidos),
                 "periodos": {k: sorted(v) for k, v in periodos.items()},
                 "sinDetalle": [os.path.relpath(r["ruta"], args.root) for r in elegidos if r["personas"] is None],
                 "erroresLectura": len(errores)})
    files = [{"fichero": os.path.relpath(r["ruta"], args.root), "company": r["company"], "period": r["period"],
              "coste": round(r["total"]["cost"], 2), "aviso": r["aviso_detalle"]} for r in elegidos]
    json.dump({"metadata": meta, "rows": filas, "files": files, "versions": versiones, "errors": errores},
              open(args.output, "w", encoding="utf-8"), ensure_ascii=False, separators=(",", ":"))
    if args.output_detail:
        json.dump({"metadata": {"read_at": meta["read_at"], "aviso": "DATOS PERSONALES: solo para la capa protegida"}, "rows": detalle},
                  open(args.output_detail, "w", encoding="utf-8"), ensure_ascii=False, separators=(",", ":"))
    print(json.dumps({"disponible": True, "ficheros": len(leidos), "periodos_elegidos": len(elegidos), "filas": len(filas),
                      "personas_mes": len(detalle), "versiones_duplicadas": len(versiones), "errores": len(errores),
                      "output": args.output}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    sys.exit(main())
